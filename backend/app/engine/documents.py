"""Files the assistant makes for the CEO: PowerPoint decks for now.

The model writes the outline (a title, slides with bullets and speaker notes); the deck is
built here with python-pptx, kept in ``settings.files_dir`` and offered for download in
the app.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from app.clock import Clock
from app.db.models import MemoRow, new_id
from app.engine.memory import Memory
from app.llm.client import LLMClient, Task
from app.llm.prompts import load_prompt
from app.llm.schemas import DeckOutline

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_SLIDES = 15
MAX_BULLETS = 5
ACCENT = RGBColor.from_string("2F5BEA")  # type: ignore[no-untyped-call]
INK = RGBColor.from_string("1F232B")  # type: ignore[no-untyped-call]
SLIDE_WIDTH = Emu(12192000)  # 16:9
SLIDE_HEIGHT = Emu(6858000)


@dataclass
class MadeFile:
    id: str
    name: str
    path: Path
    media_type: str
    slides: int = 0


def safe_name(title: str) -> str:
    words = re.sub(r"[^A-Za-z0-9 \-]", "", title).split()
    return " ".join(words[:8]) or "Presentation"


class Documents:
    def __init__(
        self,
        session_factory: sessionmaker[Session],
        llm: LLMClient,
        files_dir: Path,
        clock: Clock,
        memory: Memory | None = None,
    ) -> None:
        self._sessions = session_factory
        self._llm = llm
        self._dir = files_dir
        self._clock = clock
        self._memory = memory

    def outline(self, topic: str, brief: str, slides: int) -> DeckOutline:
        facts = self._memory.recall([topic, brief], limit=10) if self._memory else []
        context = ("Facts you may use:\n" + "\n".join(facts) + "\n\n") if facts else ""
        return self._llm.structured(
            Task.REASONING,
            "deck",
            system=[load_prompt("deck").text],
            user=(
                f"{context}Topic: {topic}\nWhat it should cover: {brief or topic}\n"
                f"Number of content slides: {slides}"
            ),
            schema=DeckOutline,
            max_tokens=6000,
        )

    def presentation(self, topic: str, brief: str = "", slides: int = 6) -> MadeFile:
        outline = self.outline(topic, brief, max(1, min(slides, MAX_SLIDES)))
        file_id = new_id()
        self._dir.mkdir(parents=True, exist_ok=True)
        path = self._dir / f"{file_id}.pptx"
        build_deck(outline, path)
        name = f"{safe_name(outline.title or topic)}.pptx"
        made = MadeFile(file_id, name, path, PPTX_TYPE, len(outline.slides[:MAX_SLIDES]))
        with self._sessions.begin() as s:
            s.add(
                MemoRow(
                    kind="file",
                    key=file_id,
                    text=name,
                    data={"path": str(path), "media_type": PPTX_TYPE, "slides": made.slides},
                    created_at=self._clock(),
                )
            )
        return made

    def get(self, file_id: str) -> MadeFile | None:
        if not re.fullmatch(r"[0-9a-f]{32}", file_id):
            return None
        with self._sessions() as s:
            row = s.scalars(
                select(MemoRow).where(MemoRow.kind == "file", MemoRow.key == file_id)
            ).first()
            if row is None:
                return None
            path = Path(str(row.data.get("path", "")))
            if not path.is_file():
                return None
            return MadeFile(file_id, row.text, path, str(row.data.get("media_type", PPTX_TYPE)))


def build_deck(outline: DeckOutline, path: Path) -> None:
    """A clean 16:9 deck: a title slide, then one slide per section with speaker notes."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT
    width, height = SLIDE_WIDTH, SLIDE_HEIGHT

    title_slide = deck.slides.add_slide(deck.slide_layouts[0])
    title_slide.shapes.title.text = outline.title
    title_slide.placeholders[1].text = outline.subtitle
    _accent_bar(title_slide, width, height)

    for item in outline.slides[:MAX_SLIDES]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = item.title
        title_frame = slide.shapes.title.text_frame
        for paragraph in title_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(32)
                run.font.bold = True
                run.font.color.rgb = INK
        body = slide.placeholders[1].text_frame
        body.clear()
        for n, bullet in enumerate(item.bullets[:MAX_BULLETS]):
            paragraph = body.paragraphs[0] if n == 0 else body.add_paragraph()
            paragraph.text = bullet.strip().rstrip(".")
            for run in paragraph.runs:
                run.font.size = Pt(22)
                run.font.color.rgb = INK
        slide.notes_slide.notes_text_frame.text = item.notes
        _accent_bar(slide, width, height)
    deck.save(str(path))


def _accent_bar(slide: object, width: int, height: int) -> None:
    shapes = slide.shapes  # type: ignore[attr-defined]
    bar = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, height - Emu(120000), width, Emu(120000))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
