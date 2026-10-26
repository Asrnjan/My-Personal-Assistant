"""Presentations, corrections that replace a proposal, and never doing the same thing twice."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from pptx import Presentation

from app.config import Settings
from app.domain.enums import ActionKind, ApprovedVia
from tests.harness import Harness, build_harness, say, tool_call
from tests.integration.test_delegation_flow import actions


@pytest.fixture
def harness(settings: Settings, tmp_path: Path):  # type: ignore[no-untyped-def]
    yield from build_harness(settings.model_copy(update={"files_dir": tmp_path / "files"}))


def test_a_deck_is_made_and_downloaded(harness: Harness) -> None:
    harness.sender.queue(
        tool_call("make_presentation", topic="Sattva Group pitch", details="Our offer", slides=2)
    )
    reply = harness.converse("Make me a deck for the Sattva Group pitch")
    result = harness.sender.calls[-1]["messages"][-1]["content"][0]["content"]
    assert '"status": "ready"' in result and '"slides": 2' in result
    [link] = reply["files"]
    assert link["name"] == "Sattva Group pitch.pptx"

    resp = harness.client.get(f"/files/{link['id']}")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "Sattva%20Group%20pitch.pptx" in resp.headers["content-disposition"]
    deck = Presentation(io.BytesIO(resp.content))
    assert len(deck.slides) == 3  # title + two content slides
    assert deck.slides[0].shapes.title.text == "Sattva Group pitch"
    assert deck.slides[1].notes_slide.notes_text_frame.text == "Say this."

    assert harness.client.get("/files/" + "0" * 32).status_code == 404
    assert harness.client.get("/files/../../etc").status_code == 404
    harness.client.headers.pop("Authorization")
    assert harness.client.get(f"/files/{link['id']}").status_code == 401


def test_a_correction_replaces_the_earlier_booking(harness: Harness) -> None:
    harness.sender.queue(
        tool_call(
            "schedule_meeting",
            title="Prep",
            date="2026-10-17",
            start_time="12:00",
            duration_minutes=30,
            attendees=[],
        ),
        say(""),
    )
    first = harness.converse("Block prep tomorrow at noon")
    assert len(first["pending_actions"]) == 1
    harness.sender.queue(
        tool_call(
            "schedule_meeting",
            title="Prep",
            date="2026-10-17",
            start_time="12:00",
            duration_minutes=60,
            attendees=[],
        ),
        say(""),
    )
    corrected = harness.converse("Make it 60 minutes")
    assert len(corrected["pending_actions"]) == 1
    assert corrected["pending_actions"][0]["id"] != first["pending_actions"][0]["id"]

    harness.converse("yes")
    assert len(harness.calendar.created) == 1  # booked once, the corrected one
    created = harness.calendar.created[0]
    assert (created.end - created.start).total_seconds() == 3600


def test_the_same_booking_never_goes_out_twice(harness: Harness) -> None:
    params = {
        "subject": "Prep",
        "start": "2026-10-17T12:00:00+05:30",
        "end": "2026-10-17T13:00:00+05:30",
        "attendees": [],
        "online": True,
    }
    one = actions(harness).propose(ActionKind.BOOK_MEETING, None, "Book. Confirm?", params)
    two = actions(harness).propose(ActionKind.BOOK_MEETING, None, "Book. Confirm?", params)
    assert actions(harness).confirm(one.id, ApprovedVia.VOICE).ok
    again = actions(harness).confirm(two.id, ApprovedVia.VOICE)
    assert again.message == "Already done a moment ago, so I didn't do it again."
    assert len(harness.calendar.created) == 1
    assert harness.audit("action.book_meeting")[-1].outcome == "duplicate"


def test_files_are_listed_and_atlas_checks_before_answering(harness: Harness) -> None:
    assert harness.client.get("/files").json() == []
    harness.sender.queue(tool_call("list_files"))
    harness.converse("Where is my deck?")
    result = harness.sender.calls[-1]["messages"][-1]["content"][0]["content"]
    assert "You haven't made any files yet." in result

    harness.sender.queue(tool_call("make_presentation", topic="Board update", slides=2))
    harness.converse("Make a board update deck")
    [listed] = harness.client.get("/files").json()
    assert listed["name"] == "Board update.pptx" and listed["slides"] == 2

    harness.sender.queue(tool_call("list_files"))
    reply = harness.converse("Show me where the deck is")
    assert reply["files"] == [{"id": listed["id"], "name": "Board update.pptx"}]
    result = harness.sender.calls[-1]["messages"][-1]["content"][0]["content"]
    assert "Board update.pptx (2 slides)" in result
